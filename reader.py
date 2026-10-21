"""Read source workbooks while retaining raw values and cell display formats."""

from __future__ import annotations

import re
import os
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

import openpyxl
import xlrd

from settings import Settings, reader_worker_count


SUPPORTED_SUFFIXES = {".xls", ".xlsx"}
TEMPLATE_FILENAMES = {"maubaocao.xls", "maubaocao.xlsx"}
RESULT_FILENAMES = {"canh_bao_du_lieu.xlsx"}
REQUIRED_HEADERS = {
    "B": {"ho dem"},
    "C": {"ten"},
    "N": {"so phieu"},
}
ZERO_FORMAT = re.compile(r"0+")


@dataclass(frozen=True)
class SourceLocation:
    file: str
    sheet: str
    row: int


@dataclass(frozen=True)
class RawRecord:
    values: Mapping[str, Any]
    display_values: Mapping[str, str]
    source: SourceLocation
    headers: Mapping[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class FileReadError:
    file: Path
    message: str


@dataclass(frozen=True)
class ReadResult:
    records: tuple[RawRecord, ...]
    files_read: tuple[Path, ...]
    errors: tuple[FileReadError, ...]
    skipped_files: tuple[Path, ...]


def read_sources(
    paths: Iterable[str | Path], settings: Settings | None = None,
    *, progress: Callable[[int, int, str], None] | None = None,
) -> ReadResult:
    """Read one or more workbooks/directories; report failures per file."""
    effective_settings = settings or Settings()
    candidates: list[Path] = []
    skipped: list[Path] = []
    errors: list[FileReadError] = []

    for input_path in paths:
        path = Path(input_path)
        if not path.exists():
            errors.append(FileReadError(path, "Đường dẫn không tồn tại."))
            continue
        if path.is_dir():
            try:
                entries = _folder_files(path, effective_settings, skipped, errors)
                for entry in entries:
                    if entry.suffix.casefold() in SUPPORTED_SUFFIXES:
                        (skipped if _excluded(entry, effective_settings) else candidates).append(entry)
            except OSError as error:
                errors.append(FileReadError(path, str(error)))
                continue
        elif path.is_file() and path.suffix.casefold() in SUPPORTED_SUFFIXES:
            (skipped if _excluded(path, effective_settings) else candidates).append(path)
        elif path.is_file():
            skipped.append(path)
        else:
            errors.append(FileReadError(path, "Đường dẫn không phải tệp hoặc thư mục."))

    unique_candidates: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        identity = str(candidate.resolve()).casefold()
        if identity not in seen:
            seen.add(identity)
            unique_candidates.append(candidate)

    # Submit only a bounded number of workbooks, preserving deterministic input
    # order in the result regardless of completion order. No Excel COM in workers.
    results = [None] * len(unique_candidates)
    workers = min(reader_worker_count(effective_settings.reader.max_workers), len(unique_candidates))
    if progress:
        progress(0, len(unique_candidates), f"Tìm thấy {len(unique_candidates)} tệp nguồn; dùng {workers} luồng đọc.")
    if workers:
        with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="pcgd-reader") as executor:
            indexes = iter(range(len(unique_candidates)))
            pending = {}
            for _ in range(workers):
                index = next(indexes)
                pending[executor.submit(_read_outcome, unique_candidates[index], effective_settings)] = index
            completed = 0
            while pending:
                finished, _ = wait(pending, return_when=FIRST_COMPLETED)
                for future in finished:
                    index = pending.pop(future)
                    results[index] = future.result()
                    completed += 1
                    if progress:
                        progress(completed, len(unique_candidates), f"Đã đọc {completed}/{len(unique_candidates)}: {unique_candidates[index].name}")
                    next_index = next(indexes, None)
                    if next_index is not None:
                        pending[executor.submit(_read_outcome, unique_candidates[next_index], effective_settings)] = next_index
    records: list[RawRecord] = []
    files_read: list[Path] = []
    for candidate, (file_records, error) in zip(unique_candidates, results):
        if error is not None:
            errors.append(error)
        else:
            records.extend(file_records)
            files_read.append(candidate.resolve())

    return ReadResult(
        records=tuple(records),
        files_read=tuple(files_read),
        errors=tuple(errors),
        skipped_files=tuple(skipped),
    )


def _read_outcome(path, settings):
    try:
        return _read_workbook(path, settings), None
    except Exception as error:
        return (), FileReadError(path.resolve(), f"{type(error).__name__}: {error}")


def _excluded(path, settings):
    resolved = path.resolve()
    return (_should_skip(path)
            or resolved.is_relative_to(settings.output_directory.resolve())
            or (settings.template_path is not None and resolved == settings.template_path.resolve()))


def _folder_files(root, settings, skipped, errors):
    if _excluded(root, settings):
        skipped.append(root)
        return
    def on_error(error):
        errors.append(FileReadError(Path(error.filename or root), str(error)))
    for directory, folders, filenames in os.walk(root, followlinks=False, onerror=on_error):
        allowed = []
        for name in sorted(folders, key=str.casefold):
            folder = Path(directory) / name
            if _excluded(folder, settings) or folder.is_symlink() or getattr(folder, "is_junction", lambda: False)():
                skipped.append(folder)
            else:
                allowed.append(name)
        folders[:] = allowed
        for name in sorted(filenames, key=str.casefold):
            yield Path(directory) / name


def _should_skip(path: Path) -> bool:
    if path.name.startswith("~$") or path.name.startswith("."):
        return True
    if path.suffix.casefold() in {".tmp", ".temp"}:
        return True
    if path.name.casefold() in TEMPLATE_FILENAMES | RESULT_FILENAMES:
        return True
    if path.stem.casefold().startswith("tonghop_"):
        return True
    if any(
        part.casefold() in {"output", "excel", "pdf", "templates", "venv", ".venv", ".git"}
        for part in path.parts
    ):
        return True
    return False


def _read_workbook(path: Path, settings: Settings) -> list[RawRecord]:
    if path.suffix.casefold() == ".xls":
        return _read_xls(path, settings)
    return _read_xlsx(path, settings)


def _read_xls(path: Path, settings: Settings) -> list[RawRecord]:
    workbook = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
    try:
        sheet, header_row = _find_xls_sheet(workbook, settings)
        def header_value(row, column):
            for r1, r2, c1, c2 in sheet.merged_cells:
                if r1 <= row < r2 and c1 <= column < c2:
                    row, column = r1, c1
                    break
            return sheet.cell_value(row, column)
        headers = _source_headers(sheet.ncols, header_row, header_value)
        start_row = max(settings.reader.data_start_row, header_row + 1)
        records: list[RawRecord] = []
        for row_index in range(start_row - 1, sheet.nrows):
            values: dict[str, Any] = {}
            displays: dict[str, str] = {}
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                column = _column_name(column_index + 1)
                raw_value = _xls_value(cell, workbook.datemode)
                values[column] = raw_value
                displays[column] = _display_value(
                    raw_value,
                    _xls_number_format(workbook, sheet, row_index, column_index),
                )
            if _is_data_row(values):
                records.append(
                    RawRecord(
                        values=values,
                        display_values=displays,
                        source=SourceLocation(str(path.resolve()), sheet.name, row_index + 1),
                        headers=headers,
                    )
                )
        return records
    finally:
        workbook.release_resources()


def _read_xlsx(path: Path, settings: Settings) -> list[RawRecord]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        sheet, header_row = _find_xlsx_sheet(workbook, settings)
        def header_value(row, column):
            row, column = row + 1, column + 1
            for area in sheet.merged_cells.ranges:
                if area.min_row <= row <= area.max_row and area.min_col <= column <= area.max_col:
                    row, column = area.min_row, area.min_col
                    break
            return sheet.cell(row, column).value
        headers = _source_headers(sheet.max_column, header_row, header_value)
        start_row = max(settings.reader.data_start_row, header_row + 1)
        records: list[RawRecord] = []
        for row_index, row in enumerate(
            sheet.iter_rows(min_row=start_row), start=start_row
        ):
            values: dict[str, Any] = {}
            displays: dict[str, str] = {}
            for column_index, cell in enumerate(row, start=1):
                column = _column_name(column_index)
                raw_value = cell.value
                values[column] = raw_value
                displays[column] = _display_value(raw_value, cell.number_format)
            if _is_data_row(values):
                records.append(
                    RawRecord(
                        values=values,
                        display_values=displays,
                        source=SourceLocation(str(path.resolve()), sheet.title, row_index),
                        headers=headers,
                    )
                )
        return records
    finally:
        workbook.close()


def _source_headers(column_count, header_row, value_at):
    headers = {}
    for column in range(column_count):
        parts = []
        for row in (header_row - 1, header_row):
            value = value_at(row, column)
            text = " ".join(str(value or "").split())
            if text and text not in parts:
                parts.append(text)
        headers[_column_name(column + 1)] = " / ".join(parts)
    return headers


def _find_xls_sheet(workbook: xlrd.book.Book, settings: Settings):
    preferred = [
        sheet for sheet in workbook.sheets()
        if sheet.name.casefold() == settings.reader.preferred_sheet.casefold()
    ]
    others = [sheet for sheet in workbook.sheets() if sheet not in preferred]
    for sheet in preferred + others:
        header_row = _find_header_row(
            sheet.nrows,
            settings.reader.max_header_scan_rows,
            lambda row, column: sheet.cell_value(row - 1, column - 1)
            if row <= sheet.nrows and column <= sheet.ncols else None,
            settings.reader.header_row,
        )
        if header_row:
            return sheet, header_row
    raise ValueError("Không tìm thấy sheet dữ liệu có tiêu đề Họ đệm, Tên, Số phiếu.")


def _find_xlsx_sheet(workbook, settings: Settings):
    preferred = [
        sheet for sheet in workbook.worksheets
        if sheet.title.casefold() == settings.reader.preferred_sheet.casefold()
    ]
    others = [sheet for sheet in workbook.worksheets if sheet not in preferred]
    for sheet in preferred + others:
        header_row = _find_header_row(
            sheet.max_row,
            settings.reader.max_header_scan_rows,
            lambda row, column: sheet.cell(row=row, column=column).value
            if row <= sheet.max_row and column <= sheet.max_column else None,
            settings.reader.header_row,
        )
        if header_row:
            return sheet, header_row
    raise ValueError("Không tìm thấy sheet dữ liệu có tiêu đề Họ đệm, Tên, Số phiếu.")


def _find_header_row(
    row_count: int, scan_limit: int, value_at, preferred_row: int
) -> int | None:
    rows = list(range(1, min(row_count, scan_limit) + 1))
    ordered_rows = ([preferred_row] if preferred_row in rows else []) + [
        row for row in rows if row != preferred_row
    ]
    for row in ordered_rows:
        matched = 0
        for column, expected in REQUIRED_HEADERS.items():
            if _normalize_header(value_at(row, _column_number(column))) in expected:
                matched += 1
        if matched == len(REQUIRED_HEADERS):
            return row
    return None


def _normalize_header(value: Any) -> str:
    import unicodedata

    text = unicodedata.normalize("NFD", str(value or "")).casefold()
    text = text.replace("đ", "d")
    return " ".join("".join(ch for ch in text if not unicodedata.combining(ch)).split())


def _is_data_row(values: Mapping[str, Any]) -> bool:
    return any(
        _has_value(values.get(_column_name(column)))
        for column in range(2, 52)
    )


def _has_value(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _xls_value(cell: xlrd.sheet.Cell, datemode: int) -> Any:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, f"Excel error {cell.value}")
    return cell.value


def _xls_number_format(
    workbook: xlrd.book.Book,
    sheet: xlrd.sheet.Sheet,
    row_index: int,
    column_index: int,
) -> str:
    try:
        xf_index = sheet.cell_xf_index(row_index, column_index)
        return workbook.format_map[workbook.xf_list[xf_index].format_key].format_str
    except (IndexError, KeyError, AttributeError, TypeError):
        return "General"


def _display_value(value: Any, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)) and ZERO_FORMAT.fullmatch(number_format):
        integer = int(value)
        if integer == value:
            return str(integer).zfill(len(number_format))
    return str(value)


def _column_number(column: str) -> int:
    number = 0
    for character in column.upper():
        number = number * 26 + ord(character) - ord("A") + 1
    return number


def _column_name(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result
