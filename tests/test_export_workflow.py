import hashlib
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from pypdf import PdfReader
import pymupdf

from src.export_workflow import export_excel_batch, export_pdf_batch, merge_pdfs, validate_pdf, validate_pdf_content
from src.template_mapper import load_template_mapping
from tests.test_approved_mapping import full_household
from scripts.verify_gui_batch_pdf import legible_fixture
from tests.test_excel_exporter import ROOT, TEMPLATE, MAPPING, COMMUNE


class ExportWorkflowTests(unittest.TestCase):
    def test_general_numeric_cell_matches_printed_integer_but_text_decimal_remains_exact(self):
        household = legible_fixture("NUMBER", 1)
        member = replace(household.members[0], fields={**household.members[0].fields, "class_name": 6.0})
        household = replace(household, members=(member,))
        text = "PHIẾU ĐIỀU TRA XÁC NHẬN CỦA UBND Xã Kiểm thử NUMBER Chủ Hộ 0912345678 Tổ kiểm thử, Xã Kiểm thử Họ1 Tên1 Kinh Không 6 THPT Kiểm thử THCS 2024 SC 25 5 24 6 1 2 0 X A Đến KT Khác Chủ hộ Cha mẹ 01/01/2010"
        reader = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda **kwargs: text)])
        mapping = load_template_mapping(MAPPING)
        validate_pdf_content(reader, household, mapping)
        textual = replace(household, members=(replace(member, fields={**member.fields, "class_name": "6.0"}),))
        with self.assertRaises(ValueError):
            validate_pdf_content(reader, textual, mapping)

    def test_wrapped_pdf_cells_are_read_in_content_order_without_losing_visual_member_order(self):
        household = legible_fixture("WRAP", 2)
        members = tuple(replace(m, fields={**m.fields, "graduation_year": "2020-2021"}) for m in household.members)
        household = replace(household, members=members)
        plain = "PHIẾU ĐIỀU TRA XÁC NHẬN CỦA UBND Xã Kiểm thử WRAP Chủ Hộ 0912345678 Tổ kiểm thử, Xã Kiểm thử Họ1 Tên1 Họ2 Tên2 Kinh Không 10A1 THPT Kiểm thử THCS 2020-2021 SC 25 5 24 6 1 2 0 X A Đến KT Khác Chủ hộ Thành viên Cha mẹ 01/01/2010 02/02/2011"
        # Layout extraction interleaves the next line of an adjacent cell.
        layout = plain.replace("2020-2021", "2020- 2026-2027 2021")
        def reader(visual, content):
            return SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda **kwargs: visual if kwargs.get("extraction_mode") == "layout" else content)])
        mapping = load_template_mapping(MAPPING)
        validate_pdf_content(reader(layout, plain), household, mapping)
        with self.assertRaises(ValueError):
            validate_pdf_content(reader(layout, plain.replace("2020-2021", "")), household, mapping)
        with self.assertRaises(ValueError):
            validate_pdf_content(reader(layout.replace("Họ1 Tên1 Họ2 Tên2", "Họ2 Tên2 Họ1 Tên1"), plain), household, mapping)

    def test_pdf_text_guard_rejects_missing_text_and_reordered_members(self):
        household = legible_fixture("CHECK", 2)
        text = "PHIẾU ĐIỀU TRA XÁC NHẬN CỦA UBND Xã Kiểm thử CHECK Chủ Hộ 0912345678 Tổ kiểm thử, Xã Kiểm thử Họ1 Tên1 Họ2 Tên2 Kinh Không 10A1 THPT Kiểm thử THCS 2024 SC 25 5 24 6 1 2 0 X A Đến KT Khác Chủ hộ Thành viên Cha mẹ 01/01/2010 02/02/2011"
        def reader(value):
            return SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda **kwargs: value)])
        mapping = load_template_mapping(MAPPING)
        validate_pdf_content(reader(text), household, mapping)
        for damaged in (text.replace("Họ2 Tên2", ""), text.replace("Họ1 Tên1 Họ2 Tên2", "Họ2 Tên2 Họ1 Tên1"), text.replace("XÁC NHẬN CỦA UBND", ""), text.replace("THPT Kiểm thử", "")):
            with self.subTest(damaged=damaged), self.assertRaises(ValueError):
                validate_pdf_content(reader(damaged), household, mapping)

    def test_exporter_generated_code_for_missing_household_code_remains_accepted(self):
        household = full_household("TEMP", 1)
        household = replace(household, household_code=None)
        with tempfile.TemporaryDirectory() as directory:
            batch = export_excel_batch([household], TEMPLATE, directory, MAPPING, commune_name=COMMUNE)
            self.assertEqual(batch.report.success_count, 1, batch.report.failures)
            self.assertEqual(len(batch.accepted), 1)
            self.assertEqual(batch.accepted[0].household, household)
            self.assertTrue(batch.accepted[0].path.name.startswith("ho_THIEU_"))
            manifest = json.loads((batch.directory / "EXPORT_REPORT.json").read_text(encoding="utf-8"))
            self.assertTrue(manifest["files"][0]["household_code"].startswith("THIEU_"))
    def test_real_com_batch_pdf_order_pages_tampering_and_no_overwrite(self):
        protected = [TEMPLATE, ROOT / "dulieumau.xls", MAPPING]
        before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in protected]
        with tempfile.TemporaryDirectory() as directory:
            households = [legible_fixture(f"H{count}", count) for count in (21, 15, 1, 8, 14, 7)]
            progress = []
            excel = export_excel_batch(households, TEMPLATE, directory, MAPPING, commune_name=COMMUNE, progress=lambda d, t, m: progress.append((d, t)))
            self.assertEqual(excel.report.failure_count, 0, excel.report.failures)
            self.assertEqual([a.household.household_code for a in excel.accepted], ["H1", "H7", "H8", "H14", "H15", "H21"])
            self.assertTrue(progress)
            self.assertEqual(progress[-1], (6, 6))
            for accepted, page_count in zip(excel.accepted, (1, 1, 2, 2, 3, 3), strict=True):
                workbook = load_workbook(accepted.path)
                try:
                    self.assertEqual(len(workbook.worksheets), page_count)
                    observed = []
                    for sheet in workbook:
                        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
                        self.assertEqual(str(sheet.page_setup.paperSize), "9")
                        self.assertEqual(sheet.page_setup.orientation, "landscape")
                        self.assertTrue(sheet.row_dimensions[44].hidden)
                        self.assertFalse(sheet.row_dimensions[43].hidden)
                        self.assertFalse(sheet.row_dimensions[82].hidden)
                        self.assertEqual(len(sheet.merged_cells.ranges), 442)
                        for row in (9, 14, 19, 24, 29, 34, 39):
                            self.assertFalse(sheet.row_dimensions[row].hidden)
                            if sheet[f"B{row}"].value:
                                observed.append((sheet[f"A{row}"].value, sheet[f"B{row}"].value))
                    self.assertEqual(observed, [(i, m.member_name) for i, m in enumerate(accepted.household.members, 1)])
                finally:
                    workbook.close()
            pdf = export_pdf_batch(excel, directory)
            self.assertEqual(pdf.failures, ())
            self.assertIsNone(pdf.merge_error)
            self.assertEqual(pdf.total_pages, 12)
            self.assertEqual([p.pages for p in pdf.successes], [1, 1, 2, 2, 3, 3])
            merged = PdfReader(pdf.merged_path)
            cursor = 0
            grid_widths = []
            for item in pdf.successes:
                reader = validate_pdf(item.path, item.pages)
                for page in reader.pages:
                    self.assertEqual(page.extract_text(), merged.pages[cursor].extract_text())
                    self.assertIn(item.household_code, page.extract_text())
                    self.assertIn("☒", page.extract_text())
                    self.assertIn("☐", page.extract_text())
                    cursor += 1
                with pymupdf.open(item.path) as rendered:
                    for page in rendered:
                        bounds = pymupdf.Rect()
                        for drawing in page.get_drawings():
                            bounds |= drawing["rect"]
                        self.assertGreater(bounds.width / (page.rect.width - 2 * 72 * 0.5 / 2.54), 0.75)
                        self.assertLess(abs((bounds.x0 + bounds.x1) / 2 - page.rect.width / 2), 2)
                        grid_widths.append(bounds.width)
            self.assertLess(max(grid_widths) - min(grid_widths), 2)
            original_pdf = pdf.merged_path.read_bytes()
            with self.assertRaises(FileExistsError):
                merge_pdfs(pdf.successes, pdf.merged_path)
            self.assertEqual(pdf.merged_path.read_bytes(), original_pdf)
            self.assertFalse(list(pdf.merged_path.parent.glob(".*.pdf")))

            damaged = excel.accepted[0].path
            workbook = load_workbook(damaged)
            workbook.worksheets[0]["B9"] = "CHANGED"
            workbook.save(damaged)
            workbook.close()
            retry = export_pdf_batch(excel, directory)
            self.assertEqual(len(retry.failures), 1)
            self.assertEqual(retry.failures[0].household_code, "H1")
            self.assertEqual(len(retry.successes), 5)
            self.assertEqual(retry.total_pages, 11)
            self.assertNotEqual(retry.merged_path, pdf.merged_path)
            self.assertEqual(pdf.merged_path.read_bytes(), original_pdf)
        self.assertEqual([hashlib.sha256(p.read_bytes()).hexdigest() for p in protected], before)

    def test_duplicate_codes_are_reported_without_using_one_households_data_for_another(self):
        with tempfile.TemporaryDirectory() as directory:
            report = export_excel_batch([full_household("DUP", 1), full_household("DUP", 7)], TEMPLATE, directory, MAPPING, commune_name=COMMUNE)
            self.assertEqual(report.report.success_count, 0)
            self.assertEqual(report.report.failure_count, 2)
            self.assertEqual(report.accepted, ())
            self.assertFalse(list(report.directory.glob("*.xlsx")))
            with self.assertRaises(ValueError):
                export_pdf_batch(report, directory)
