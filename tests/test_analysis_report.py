import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from analysis import analyze_sources
from settings import Settings, load_settings, save_settings, validate_export_settings
from warning_report import write_warning_report


class AnalysisReportTests(unittest.TestCase):
    def test_settings_load_unconfigured_and_configured_commune_from_temporary_files(self):
        for commune_name, expected_commune in (
            (None, ""),
            ("", ""),
            ("   ", ""),
            ("PDP", "PDP"),
            (" PDP ", "PDP"),
            ("Xã Kiểm thử", "Xã Kiểm thử"),
        ):
            with self.subTest(commune_name=commune_name):
                with tempfile.TemporaryDirectory() as directory:
                    settings_path = Path(directory) / "settings.json"
                    settings_path.write_text(
                        json.dumps(
                            {
                                "commune_name": commune_name,
                                "school_year": "2026-2027",
                            },
                            ensure_ascii=False,
                        ),
                        encoding="utf-8",
                    )

                    settings = load_settings(settings_path)

                    self.assertEqual(settings.commune_name, expected_commune)
                    self.assertEqual(settings.school_year, "2026-2027")

    def test_export_settings_accept_consecutive_school_years_and_reject_invalid_values(self):
        for school_year in ("2026-2027", "2027-2028"):
            with self.subTest(school_year=school_year):
                validate_export_settings(
                    Settings(commune_name="Xã kiểm thử", school_year=school_year)
                )

        for school_year in ("2026/2027", "2026-2028", "năm học"):
            with self.subTest(school_year=school_year):
                with self.assertRaisesRegex(ValueError, "school_year"):
                    validate_export_settings(
                        Settings(commune_name="Xã kiểm thử", school_year=school_year)
                    )

    def test_save_settings_updates_locality_and_school_year_without_dropping_other_keys(self):
        with tempfile.TemporaryDirectory() as directory:
            settings_path = Path(directory) / "settings.json"
            original = {
                "commune_name": None,
                "school_year": None,
                "reader": {"preferred_sheet": "CustomSheet"},
                "source_columns": {"phone": "AW"},
                "output_directory": "custom-output",
                "source_path": "source.xls",
                "template_path": "template.xls",
                "other_setting": {"preserved": True},
            }
            settings_path.write_text(
                json.dumps(original),
                encoding="utf-8",
            )

            settings = save_settings(
                settings_path,
                commune_name="  ",
                school_year=" 2026-2027 ",
            )
            saved = json.loads(settings_path.read_text(encoding="utf-8"))

            self.assertIsNone(saved["commune_name"])
            self.assertEqual(saved["school_year"], "2026-2027")
            self.assertEqual(saved["reader"], original["reader"])
            self.assertEqual(saved["source_columns"], original["source_columns"])
            self.assertEqual(saved["output_directory"], original["output_directory"])
            self.assertEqual(saved["source_path"], original["source_path"])
            self.assertEqual(saved["template_path"], original["template_path"])
            self.assertEqual(saved["other_setting"], original["other_setting"])
            self.assertEqual(settings.commune_name, "")
            self.assertEqual(settings.school_year, "2026-2027")
            with self.assertRaisesRegex(ValueError, "commune_name"):
                validate_export_settings(settings)

            saved_settings = save_settings(
                settings_path,
                commune_name=" Xã Kiểm Thử ",
                school_year="2026-2027",
                source_path="selected-source.xls",
                template_path="selected-template.xls",
                output_directory="selected-output",
            )
            self.assertEqual(saved_settings.commune_name, "Xã Kiểm Thử")
            self.assertEqual(saved_settings.school_year, "2026-2027")
            self.assertEqual(
                saved_settings.source_path,
                settings_path.parent.parent / "selected-source.xls",
            )
            self.assertEqual(
                saved_settings.template_path,
                settings_path.parent.parent / "selected-template.xls",
            )
            self.assertEqual(
                saved_settings.output_directory,
                settings_path.parent.parent / "selected-output",
            )
            validate_export_settings(saved_settings)

    def test_verified_fixture_analysis_and_warning_report(self):
        fixture = Path(__file__).resolve().parents[1] / "dulieumau.xls"
        if not fixture.exists():
            self.skipTest("Verified source workbook is not present")
        with tempfile.TemporaryDirectory() as directory:
            report_path = Path(directory) / "CANH_BAO_DU_LIEU.xlsx"
            result = analyze_sources([fixture])
            write_warning_report(result, report_path)

            self.assertEqual(len(result.read_result.files_read), 1)
            self.assertEqual(len(result.members), 427)
            self.assertEqual(len(result.grouping_result.households), 120)
            self.assertTrue(report_path.exists())
            workbook = load_workbook(report_path, read_only=True, data_only=True)
            try:
                summary = dict(workbook["Tong quan"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(summary["Tổng số tệp đọc"], 1)
                self.assertEqual(summary["Tổng số bản ghi"], 427)
                self.assertEqual(summary["Tổng số hộ"], 120)
                household_rows = list(workbook["So ho"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(household_rows), 120)
                self.assertEqual(sum(row[3] for row in household_rows), 427)
                warnings = list(workbook["Canh bao"].iter_rows(min_row=2, values_only=True))
                codes = [row[1] for row in warnings]
                self.assertEqual(codes.count("HOUSEHOLD_HEAD_NAME_CONFLICT"), 3)
                self.assertEqual(codes.count("HOUSEHOLD_RESIDENCE_CONFLICT"), 5)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
