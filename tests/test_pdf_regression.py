"""Regression tests at the public Excel -> PDF batch boundary, real COM."""
import tempfile
import unittest
from dataclasses import replace
from copy import copy

from pypdf import PdfReader
from pypdf import PdfWriter
from openpyxl import load_workbook
from openpyxl.styles import Side
from pathlib import Path
from unittest.mock import patch
from src.export_workflow import export_excel_batch, export_pdf_batch, validate_pdf, fingerprint
from scripts.verify_gui_batch_pdf import legible_fixture
from tests.test_excel_exporter import TEMPLATE, MAPPING, COMMUNE


class PdfPaginationRegressionTests(unittest.TestCase):
    def test_changed_border_is_rejected_even_if_input_hash_is_resealed(self):
        with tempfile.TemporaryDirectory() as root:
            excel = export_excel_batch([legible_fixture("BORDER", 1)], TEMPLATE, root, MAPPING, commune_name=COMMUNE)
            self.assertEqual(excel.report.failures, ())
            accepted = excel.accepted[0]
            book = load_workbook(accepted.path)
            try:
                border = copy(book.active["B9"].border)
                border.left = Side(style="thin", color="FFFF0000")
                book.active["B9"].border = border
                book.save(accepted.path)
            finally:
                book.close()
            # Exercise structural acceptance independently of the earlier hash
            # guard: the modified file must still fail its border comparison.
            resealed = replace(accepted, sha256=fingerprint(accepted.path))
            pdf = export_pdf_batch(replace(excel, accepted=(resealed,)), root)
            self.assertEqual(pdf.successes, ())
            self.assertEqual(len(pdf.failures), 1)
            self.assertIn("formatting differs", pdf.failures[0].error)

    def test_long_class_uses_smaller_local_font_and_remains_complete_in_pdf(self):
        household = legible_fixture("LONG-CLASS", 1)
        member = household.members[0]
        member = replace(member, fields={**member.fields, "class_name": "11 CHUYÊN TRUNG"})
        household = replace(household, members=(member,))
        with tempfile.TemporaryDirectory() as root:
            excel = export_excel_batch([household], TEMPLATE, root, MAPPING, commune_name=COMMUNE)
            self.assertEqual(excel.report.failures, ())
            workbook = load_workbook(excel.accepted[0].path)
            try:
                self.assertLess(workbook.active["H9"].font.sz, workbook.active["H14"].font.sz)
                self.assertGreaterEqual(workbook.active["H9"].font.sz, 8)
                self.assertEqual(workbook.active["H9"].value, "11 CHUYÊN TRUNG")
            finally:
                workbook.close()
            pdf = export_pdf_batch(excel, root)
            self.assertEqual(pdf.failures, ())
            self.assertEqual(pdf.total_pages, 1)

    def test_wrong_number_of_pages_is_still_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "two.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=841.89, height=595.28)
            writer.add_blank_page(width=841.89, height=595.28)
            writer.write(path)
            writer.close()
            with self.assertRaisesRegex(ValueError, "page count"):
                validate_pdf(path, 1)

    def test_com_export_failure_does_not_poison_next_household(self):
        import win32com.client
        import pywintypes
        # Fault injection only at the external COM boundary; all remaining
        # workbook operations and the following household use real Excel.
        original_dispatch = win32com.client.DispatchEx
        state = {"failed": False}
        class WorkbookProxy:
            def __init__(self, real):
                self.real = real
            def __getattr__(self, name):
                return getattr(self.real, name)
            def ExportAsFixedFormat(self, *args, **kwargs):
                if not state["failed"]:
                    state["failed"] = True
                    raise pywintypes.com_error(-2147352567, "Injected ExportAsFixedFormat failure", None, None)
                return self.real.ExportAsFixedFormat(*args, **kwargs)
        class WorkbooksProxy:
            def __init__(self, real):
                self.real = real
            def Open(self, *args, **kwargs):
                return WorkbookProxy(self.real.Open(*args, **kwargs))
        class ExcelProxy:
            def __init__(self, real):
                object.__setattr__(self, "real", real)
            def __getattr__(self, name):
                if name == "Workbooks":
                    return WorkbooksProxy(self.real.Workbooks)
                return getattr(self.real, name)
            def __setattr__(self, name, value):
                setattr(self.real, name, value)
        with tempfile.TemporaryDirectory() as root:
            excel = export_excel_batch([legible_fixture("FAIL-1", 1), legible_fixture("PASS-2", 1)], TEMPLATE, root, MAPPING, commune_name=COMMUNE)
            self.assertEqual(excel.report.failure_count, 0, excel.report.failures)
            with patch("win32com.client.DispatchEx", side_effect=lambda *a, **k: ExcelProxy(original_dispatch(*a, **k))):
                pdf = export_pdf_batch(excel, root)
            self.assertEqual([f.household_code for f in pdf.failures], ["FAIL-1"])
            self.assertIn("Injected ExportAsFixedFormat failure", pdf.failures[0].error)
            self.assertEqual([p.household_code for p in pdf.successes], ["PASS-2"])
            self.assertEqual(pdf.total_pages, 1)
            self.assertEqual(len(PdfReader(pdf.merged_path).pages), 1)

    def test_seven_members_share_one_sheet_and_one_page(self):
        with tempfile.TemporaryDirectory() as root:
            excel = export_excel_batch([legible_fixture("SEVEN", 7)], TEMPLATE, root, MAPPING, commune_name=COMMUNE)
            self.assertEqual(excel.report.failure_count, 0, excel.report.failures)
            book = load_workbook(excel.accepted[0].path)
            try:
                self.assertEqual(len(book.worksheets), 1)
                self.assertEqual(book.worksheets[0]["A39"].value, 7)
                self.assertFalse(book.worksheets[0].row_dimensions[39].hidden)
            finally:
                book.close()
            pdf = export_pdf_batch(excel, root)
            self.assertEqual(pdf.failures, (), pdf.failures)
            self.assertEqual(pdf.total_pages, 1)

    def test_single_household_keeps_footer_on_one_a4_page(self):
        with tempfile.TemporaryDirectory() as root:
            excel = export_excel_batch([legible_fixture("PDF-REGRESSION", 1)], TEMPLATE, root, MAPPING, commune_name=COMMUNE)
            self.assertEqual(excel.report.failure_count, 0, excel.report.failures)
            pdf = export_pdf_batch(excel, root)
            self.assertEqual(pdf.failures, (), pdf.failures)
            self.assertEqual(pdf.total_pages, 1)
            text = PdfReader(pdf.successes[0].path).pages[0].extract_text(extraction_mode="layout")
            self.assertIn("Chủ hộ gia đình", text)


if __name__ == "__main__":
    unittest.main()
