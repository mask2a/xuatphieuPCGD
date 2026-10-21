import tempfile
import unittest
import json
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from scripts.verify_gui_batch_pdf import legible_fixture
from src.export_workflow import export_excel_batch, export_pdf_batch
from tests.test_excel_exporter import TEMPLATE, MAPPING, COMMUNE


class PdfOptionTests(unittest.TestCase):
    def test_unchecked_applies_to_each_source_and_preserves_multiple_pages(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            households = [replace(legible_fixture("FAST-8", 8), source_file=str(root / "nguon A.xlsx")), replace(legible_fixture("FAST-1", 1), source_file=str(root / "nguon B.xlsx"))]
            batch = export_excel_batch(households, TEMPLATE, root, MAPPING, commune_name=COMMUNE)
            self.assertEqual(batch.report.failures, ())
            for item in batch.accepted:
                workbook = load_workbook(item.path)
                workbook.worksheets[0]["B9"] = "Dữ liệu đã sửa"
                workbook.save(item.path)
                workbook.close()
            pdf = export_pdf_batch(batch, root, validate=False)
            self.assertEqual(pdf.failures, ())
            self.assertEqual(pdf.total_pages, 3)
            self.assertEqual({p.name for p in pdf.merged_paths}, {"nguon A.pdf", "nguon B.pdf"})
            self.assertEqual(len(PdfReader(root / "nguon A.pdf").pages), 2)
            self.assertEqual(len(PdfReader(root / "nguon B.pdf").pages), 1)
            for item in pdf.successes:
                report = json.loads((item.path.parent / "PDF_REPORT.json").read_text(encoding="utf-8"))
                self.assertIs(report["validation_enabled"], False)

    def test_unchecked_exports_current_excel_while_checked_rejects_changed_excel(self):
        with tempfile.TemporaryDirectory() as temp:
            batch = export_excel_batch([legible_fixture("OPTION", 1)], TEMPLATE, temp, MAPPING, commune_name=COMMUNE)
            self.assertEqual(batch.report.failures, ())
            workbook = load_workbook(batch.accepted[0].path)
            workbook.active["B9"] = "Tên đã sửa"
            workbook.save(batch.accepted[0].path)
            workbook.close()
            checked = export_pdf_batch(batch, temp)
            self.assertEqual(len(checked.failures), 1)
            self.assertIsNone(checked.merged_path)
            without_template = replace(batch, template=Path(temp) / "missing.xls", mapping=Path(temp) / "missing.json")
            unchecked = export_pdf_batch(without_template, temp, validate=False)
            self.assertEqual(unchecked.failures, ())
            self.assertEqual(len(unchecked.successes), 1)
            self.assertIn("Tên đã sửa", PdfReader(unchecked.merged_path).pages[0].extract_text())
