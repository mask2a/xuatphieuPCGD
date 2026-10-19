import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from normalizer import normalize_records
from reader import read_sources
from settings import Settings


class NormalizerTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.source = Path(self.temp_dir.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MauNhapLieu"
        for column, value in {
            1: "TT", 2: "Họ đệm", 3: "Tên", 4: "Ngày", 5: "Tháng",
            6: "Năm sinh", 11: "Chủ hộ", 12: "Tên", 14: "Số phiếu",
            15: "Diện cư trú", 47: "Quan hệ với chủ hộ", 49: "Điện thoại",
        }.items():
            sheet.cell(row=2, column=column, value=value)
        for column, value in {
            2: "  Nguyễn   Van ", 3: "  an  ", 4: 1, 5: 2, 14: " 001 ",
            11: "Nguyễn", 12: "An", 15: "Thường trú", 47: "Chủ hộ",
            49: " 001234 ",
        }.items():
            sheet.cell(row=5, column=column, value=value)
        sheet.cell(row=5, column=4).number_format = "00"
        sheet.cell(row=5, column=5).number_format = "00"
        workbook.save(self.source)
        workbook.close()

    def test_normalizes_names_and_codes_without_mutating_source_record(self):
        raw = read_sources([self.source]).records[0]

        member = normalize_records([raw], Settings(commune_name="Phường X"))[0]

        self.assertEqual(member.member_name, "Nguyễn Van an")
        self.assertEqual(member.member_name_key, "nguyễn van an")
        self.assertEqual(member.household_code, " 001 ")
        self.assertEqual(member.household_code_key, "001")
        self.assertEqual(member.birth_components, ("01", "02", None))
        self.assertEqual(member.phone, " 001234 ")
        self.assertEqual(member.commune_name, "Phường X")
        self.assertEqual(raw.values["B"], "  Nguyễn   Van ")
        self.assertEqual(raw.values["N"], " 001 ")

    def test_blank_values_remain_blank_instead_of_being_inferred(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MauNhapLieu"
        for column, value in {1: "TT", 2: "Họ đệm", 3: "Tên", 14: "Số phiếu"}.items():
            sheet.cell(row=2, column=column, value=value)
        sheet.cell(row=5, column=2, value="Tên")
        sheet.cell(row=5, column=14, value="A")
        workbook.save(self.source)
        workbook.close()

        member = normalize_records(read_sources([self.source]).records)[0]

        self.assertEqual(member.birth_components, (None, None, None))
        self.assertIsNone(member.phone)
        self.assertEqual(member.member_name, "Tên")


if __name__ == "__main__":
    unittest.main()
