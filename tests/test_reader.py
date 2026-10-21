import tempfile
import unittest
from pathlib import Path
from dataclasses import replace
from threading import Barrier
from unittest.mock import patch

from openpyxl import Workbook

from reader import read_sources
from reader import _read_workbook
from settings import Settings


class ReaderTests(unittest.TestCase):
    def test_parallel_read_keeps_input_order_and_deduplicates_paths(self):
        first, second = self.root / "z.xlsx", self.root / "a.xlsx"
        self.create_source_workbook(first)
        self.create_source_workbook(second)
        barrier = Barrier(2, timeout=5)
        def overlapping_read(path, settings):
            barrier.wait()
            return _read_workbook(path, settings)
        settings = replace(self.settings, reader=replace(self.settings.reader, max_workers=2))
        progress = []
        with patch("reader._read_workbook", side_effect=overlapping_read):
            result = read_sources([first, second, first], settings, progress=lambda done, total, message: progress.append((done, total)))
        self.assertEqual(result.errors, ())
        self.assertEqual(result.files_read, (first.resolve(), second.resolve()))
        self.assertEqual([r.source.file for r in result.records], [str(first.resolve()), str(second.resolve())])
        self.assertEqual(progress[-1], (2, 2))

    def test_folder_read_excludes_custom_output_and_custom_template(self):
        output = self.root / "ket_qua"
        output.mkdir()
        self.create_source_workbook(output / "generated.xlsx")
        template = self.root / "custom_template.xlsx"
        self.create_source_workbook(template)
        source = self.root / "source.xlsx"
        self.create_source_workbook(source)
        settings = replace(self.settings, output_directory=output, template_path=template)
        result = read_sources([self.root], settings)
        self.assertEqual(result.files_read, (source.resolve(),))
        self.assertEqual(result.errors, ())

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.settings = Settings()

    def create_source_workbook(self, path: Path, phone_format: str = "General"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MauNhapLieu"
        sheet.append(["Thông tin đối tượng"])
        headers = ["TT", "Họ đệm", "Tên", "Ngày", "Tháng", "Năm sinh"]
        for column, value in zip((1, 2, 3, 4, 5, 6), headers):
            sheet.cell(row=2, column=column, value=value)
        sheet.cell(row=2, column=14, value="Số phiếu")
        sheet.cell(row=2, column=49, value="Điện thoại")
        sheet.cell(row=5, column=2, value="  Nguyễn   Văn ")
        sheet.cell(row=5, column=3, value="An")
        sheet.cell(row=5, column=4, value=1)
        sheet.cell(row=5, column=4).number_format = "00"
        sheet.cell(row=5, column=5, value=2)
        sheet.cell(row=5, column=5).number_format = "00"
        sheet.cell(row=5, column=6, value=None)
        sheet.cell(row=5, column=14, value="001")
        sheet.cell(row=5, column=49, value=946236722)
        sheet.cell(row=5, column=49).number_format = phone_format
        workbook.save(path)
        workbook.close()

    def test_reads_xlsx_source_and_preserves_displayed_leading_zeroes(self):
        source = self.root / "source.xlsx"
        self.create_source_workbook(source, "0000000000")

        result = read_sources([source], self.settings)

        self.assertEqual(len(result.records), 1)
        record = result.records[0]
        self.assertEqual(record.values["N"], "001")
        self.assertEqual(record.display_values["D"], "01")
        self.assertEqual(record.display_values["E"], "02")
        self.assertEqual(record.values["F"], None)
        self.assertEqual(record.display_values["AW"], "0946236722")
        self.assertEqual(record.source.sheet, "MauNhapLieu")
        self.assertEqual(record.source.row, 5)
        self.assertEqual(result.files_read, (source.resolve(),))
        self.assertEqual(result.errors, ())

    def test_folder_read_skips_templates_temporary_and_output_files(self):
        source = self.root / "source.xlsx"
        self.create_source_workbook(source)
        self.create_source_workbook(self.root / "maubaocao.xlsx")
        self.create_source_workbook(self.root / "~$locked.xlsx")
        output = self.root / "output"
        output.mkdir()
        self.create_source_workbook(output / "generated.xlsx")

        result = read_sources([self.root], self.settings)

        self.assertEqual(len(result.records), 1)
        self.assertEqual(result.records[0].source.file, str(source.resolve()))
        self.assertEqual(result.files_read, (source.resolve(),))
        self.assertEqual(len(result.skipped_files), 3)

    def test_one_bad_file_does_not_block_other_files(self):
        valid = self.root / "valid.xlsx"
        self.create_source_workbook(valid)
        invalid = self.root / "broken.xlsx"
        invalid.write_bytes(b"not an Excel workbook")

        result = read_sources([valid, invalid], self.settings)

        self.assertEqual(len(result.records), 1)
        self.assertEqual(result.files_read, (valid.resolve(),))
        self.assertEqual(len(result.errors), 1)
        self.assertEqual(result.errors[0].file, invalid.resolve())

    def test_reads_verified_legacy_xls_fixture(self):
        fixture = Path(__file__).resolve().parents[1] / "dulieumau.xls"
        if not fixture.exists():
            self.skipTest("Verified source workbook is not present")

        result = read_sources([fixture], self.settings)

        self.assertEqual(len(result.records), 427)
        self.assertEqual(len({record.values["N"] for record in result.records}), 120)
        self.assertEqual(result.records[0].source.row, 5)
        self.assertEqual(result.records[0].display_values["AW"], "0946236722")
        self.assertEqual(result.records[0].display_values["D"], "01")
        self.assertEqual(
            sum(not record.display_values["F"].strip() for record in result.records), 1
        )
        missing_year_record = next(
            record for record in result.records if not record.display_values["F"].strip()
        )
        self.assertEqual(missing_year_record.values["F"], " ")


if __name__ == "__main__":
    unittest.main()
