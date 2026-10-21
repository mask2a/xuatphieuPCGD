import tempfile
import unittest
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from household_grouper import group_households
from normalizer import normalize_records
from reader import RawRecord, SourceLocation
from settings import Settings
from src.excel_exporter import ExportConfigurationError, export_households
from src.school_year_labels import YEAR_LABEL_CELLS


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "tempplates/maubaocao.xls" if (ROOT / "tempplates/maubaocao.xls").is_file() else ROOT / "maubaocao.xls"
MAPPING = ROOT / "config" / "template_mapping.json"
COMMUNE = "Xã Kiểm thử"


def make_household(
    code,
    count,
    *,
    missing_birth_year_for=(),
    head_phone="0912345678",
    non_head_phones=(),
    head_names=None,
    residences=None,
    unmapped_fields=None,
    source_values_by_member=None,
    commune_name=COMMUNE,
):
    records = []
    for index in range(count):
        family_name = f"Họ{index + 1}"
        given_name = f"Tên{index + 1}"
        birth_year = None if index in missing_birth_year_for else 2010 + index
        head_name = (head_names or ["Chủ Hộ"])[min(index, len(head_names or ["Chủ Hộ"]) - 1)]
        residence = (residences or ["Thường trú"])[min(index, len(residences or ["Thường trú"]) - 1)]
        relationship = "Chủ hộ" if index == 0 else "Thành viên"
        phone = head_phone if index == 0 else (
            non_head_phones[index - 1]
            if index - 1 < len(non_head_phones)
            else None
        )
        values = {
            "B": family_name,
            "C": given_name,
            "D": index + 1,
            "E": index + 2,
            "F": birth_year,
            "G": "X" if index % 2 == 0 else None,
            "H": "Kinh",
            "I": None,
            "K": head_name,
            "L": None,
            "M": None,
            "N": code,
            "O": residence,
            "P": None,
            "Q": None,
            "R": None,
            "S": None,
            "T": None,
            "AU": relationship,
            "AV": "Cha mẹ",
            "AW": phone,
        }
        values.update(unmapped_fields or {})
        values.update((source_values_by_member or {}).get(index, {}))
        displays = {
            column: "" if value is None else str(value)
            for column, value in values.items()
        }
        displays["D"] = f"{index + 1:02}"
        displays["E"] = f"{index + 2:02}"
        records.append(
            RawRecord(
                values=values,
                display_values=displays,
                source=SourceLocation(
                    "synthetic-source.xls", "MauNhapLieu", index + 5
                ),
            )
        )
    members = normalize_records(records, Settings(commune_name=commune_name))
    return group_households(members).households[0]


def write_export_settings(directory, *, commune_name=COMMUNE, school_year="2025-2026"):
    path = Path(directory) / "settings.json"
    path.write_text(
        json.dumps(
            {"commune_name": commune_name, "school_year": school_year},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def _read_template_year_labels():
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    workbook = None
    try:
        workbook = excel.Workbooks.Open(
            str(TEMPLATE),
            UpdateLinks=0,
            ReadOnly=True,
            AddToMru=False,
            IgnoreReadOnlyRecommended=True,
        )
        sheet = workbook.Worksheets.Item(1)
        return {
            address: sheet.Range(address).Value2
            for address in YEAR_LABEL_CELLS
        }
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()
        pythoncom.CoUninitialize()


class ExcelExporterTests(unittest.TestCase):
    def test_gui_test_export_applies_configured_year_labels_to_the_copy(self):
        household = make_household("GUI-TEST", 1)
        template_hash_before = hashlib.sha256(TEMPLATE.read_bytes()).digest()

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            report = export_households(
                (household,),
                TEMPLATE,
                output,
                MAPPING,
                settings_path=write_export_settings(
                    directory,
                    school_year="2026-2027",
                ),
                apply_school_year_labels=True,
            )

            self.assertEqual(report.success_count, 1)
            self.assertEqual(report.failure_count, 0)
            workbook = load_workbook(report.successes[0].path, data_only=True)
            try:
                sheet = workbook["Phiếu điều tra 1 mặt"]
                expected = (
                    "2025-2026",
                    "2026-2027",
                    "2027-2028",
                    "2028-2029",
                    "2029-2030",
                )
                self.assertEqual(
                    tuple(sheet[address].value for address in YEAR_LABEL_CELLS[:5]),
                    expected,
                )
                for slot in range(14):
                    addresses = YEAR_LABEL_CELLS[slot * 5:(slot + 1) * 5]
                    self.assertEqual(
                        tuple(sheet[address].value for address in addresses),
                        expected,
                    )
                self.assertEqual(len(sheet.merged_cells.ranges), 442)
            finally:
                workbook.close()

        self.assertEqual(
            hashlib.sha256(TEMPLATE.read_bytes()).digest(),
            template_hash_before,
        )

    def test_exports_short_full_and_supplementary_households_from_real_template(self):
        if not TEMPLATE.exists():
            self.skipTest("The verified source template is not present")
        households = tuple(
            make_household(f"{count:02}", count)
            for count in (1, 6, 7, 14, 15)
        )
        template_hash_before = hashlib.sha256(TEMPLATE.read_bytes()).digest()
        source_path = ROOT / "dulieumau.xls"
        source_hash_before = hashlib.sha256(source_path.read_bytes()).digest()

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "excel"
            report = export_households(
                households,
                TEMPLATE,
                output,
                MAPPING,
                settings_path=write_export_settings(directory),
            )

            self.assertEqual(report.success_count, 5)
            self.assertEqual(report.failure_count, 0)
            self.assertEqual(report.summary["success_count"], 5)
            self.assertEqual(report.summary["failure_count"], 0)
            self.assertGreater(report.summary["warning_count"], 0)
            self.assertTrue(
                all(
                    "SCHOOL_YEAR_ROWS_NOT_MAPPED"
                    in {warning.code for warning in item.warnings}
                    for item in report.successes
                )
            )
            self.assertTrue(
                all(
                    any(
                        "2025-2026" in warning.message
                        for warning in item.warnings
                        if warning.code == "SCHOOL_YEAR_ROWS_NOT_MAPPED"
                    )
                    for item in report.successes
                )
            )
            self.assertEqual(
                {item.member_count for item in report.successes},
                {1, 6, 7, 14, 15},
            )
            self.assertEqual(
                hashlib.sha256(TEMPLATE.read_bytes()).digest(),
                template_hash_before,
            )
            self.assertEqual(
                hashlib.sha256(source_path.read_bytes()).digest(),
                source_hash_before,
            )

            by_code = {str(item.household_code): item for item in report.successes}
            for code, count in (("01", 1), ("06", 6), ("07", 7), ("14", 14), ("15", 15)):
                exported = by_code[code]
                self.assertTrue(exported.path.is_file())
                workbook = load_workbook(exported.path, data_only=False)
                try:
                    sheet = workbook["Phiếu điều tra 1 mặt"]
                    self.assertEqual(sheet["A1"].value, f"Xã (1): {COMMUNE}")
                    self.assertEqual(sheet["Z1"].value, f"Số phiếu (4): {code}")
                    self.assertEqual(
                        sheet["Z3"].value, "Điện thoại (6): 0912345678"
                    )
                    self.assertEqual(
                        sheet["AA82"].value.splitlines()[0],
                        "XÁC NHẬN CỦA UBND XÃ " + COMMUNE,
                    )
                    self.assertEqual(len(sheet.merged_cells.ranges), 442)
                    if count <= 6:
                        first_hidden = 47
                        self.assertTrue(sheet.row_dimensions[first_hidden].hidden)
                        self.assertEqual(sheet.print_area, "'Phiếu điều tra 1 mặt'!$A$1:$AG$87")
                        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
                    else:
                        self.assertFalse(sheet.row_dimensions[39].hidden)
                    self.assertIn("Họ1 Tên1", sheet["B9"].value)
                    self.assertEqual(sheet["D11"].value, "01/02/2010")
                    self.assertEqual(sheet["B10"].value, "QH với chủ hộ: Chủ hộ")
                    self.assertEqual(sheet["D12"].value, "DT: Kinh")
                    self.assertEqual(sheet["B13"].value, "Cha, mẹ, Ng đỡ đầu: Cha mẹ")
                    self.assertEqual(sheet["A1"].font.name, "Arial")
                    self.assertTrue(sheet["A1"].font.bold)
                    self.assertEqual(sheet["C12"].font.name, "Segoe UI Symbol")
                    self.assertTrue(sheet["B9"].border.bottom.style)
                    self.assertTrue(sheet["AA82"].value.startswith("XÁC NHẬN CỦA UBND"))
                    self.assertFalse(workbook._external_links)
                    for member_index in range(count):
                        page_index, slot_index = divmod(member_index, 7)
                        member_sheet = workbook.worksheets[page_index]
                        start_row = (9, 14, 19, 24, 29, 34, 39, 47, 52, 57, 62, 67, 72, 77)[slot_index]
                        self.assertEqual(
                            member_sheet[f"B{start_row}"].value,
                            f"Họ{member_index + 1} Tên{member_index + 1}",
                        )
                    if count == 15:
                        self.assertEqual(len(workbook.worksheets), 3)
                        supplement = workbook.worksheets[2]
                        self.assertEqual(supplement["A9"].value, 15)
                        self.assertEqual(supplement["B9"].value, "Họ15 Tên15")
                        self.assertTrue(supplement.row_dimensions[47].hidden)
                finally:
                    workbook.close()

    def test_missing_and_conflicting_values_are_not_inferred_or_silently_selected(self):
        households = (
            make_household(
                "NO-PHONE",
                2,
                missing_birth_year_for=(0,),
                head_phone=None,
                non_head_phones=("0900111222",),
            ),
            make_household(
                "CONFLICT",
                2,
                head_names=("Chủ Một", "Chủ Hai"),
                residences=("Thường trú", "Tạm trú"),
                unmapped_fields={
                    "J": "Diện ưu tiên",
                    "M": "Tổ 7",
                    "Q": "1",
                    "R": "Lớp 1",
                    "S": "Huyện thử",
                    "T": "Trường thử",
                    "V": "THCS",
                    "W": "Bổ túc",
                    "X": "2025",
                    "Y": "Nghề",
                    "AA": "2026",
                    "AB": "Đã học xong",
                    "AC": "2025",
                    "AD": "Lớp 5",
                    "AE": "2024",
                    "AF": "Lớp 6",
                    "AG": "Lớp 5",
                    "AH": "Mức 1",
                    "AX": "Ghi chú",
                },
            ),
        )

        with tempfile.TemporaryDirectory() as directory:
            report = export_households(
                households,
                TEMPLATE,
                Path(directory),
                MAPPING,
                settings_path=write_export_settings(directory),
            )

            self.assertEqual(report.success_count, 2)
            self.assertEqual(report.failure_count, 0)
            by_code = {str(item.household_code): item for item in report.successes}
            missing = load_workbook(by_code["NO-PHONE"].path, data_only=False)
            conflict = load_workbook(by_code["CONFLICT"].path, data_only=False)
            try:
                missing_sheet = missing["Phiếu điều tra 1 mặt"]
                self.assertEqual(
                    missing_sheet["Z3"].value.replace("Điện thoại (6):", "").strip("_ "),
                    "",
                )
                self.assertEqual(missing_sheet["D11"].value, "01/02/")
                self.assertNotIn(
                    "HOUSEHOLD_PHONE_CONFLICT",
                    {warning.code for warning in by_code["NO-PHONE"].warnings},
                )
                conflict_sheet = conflict["Phiếu điều tra 1 mặt"]
                self.assertTrue(conflict_sheet["H3"].value.endswith("_" * 10))
                self.assertNotIn("X", conflict_sheet["Z2"].value)
                self.assertEqual(conflict_sheet["Z2"].value.count("☒"), 0)
                self.assertEqual(conflict_sheet["Z2"].value.count("☐"), 2)
                self.assertIn("HOUSEHOLD_RESIDENCE_NOT_EXPORTED", {w.code for w in by_code["CONFLICT"].warnings})
                conflict_codes = {
                    warning.code for warning in by_code["CONFLICT"].warnings
                }
                self.assertIn("HOUSEHOLD_HEAD_NAME_CONFLICT", conflict_codes)
                self.assertIn("HOUSEHOLD_RESIDENCE_CONFLICT", conflict_codes)
                self.assertNotIn("UNMAPPED_ADDRESS", conflict_codes)
                self.assertIn("B5_GRADE_CLASS_NOT_COMPARABLE", conflict_codes)
                self.assertNotIn("UNMAPPED_NOTES", conflict_codes)
                self.assertIn("UNMAPPED_PRIORITY_CATEGORY", conflict_codes)
                self.assertEqual(conflict_sheet["H9"].value, "Lớp 1")
                self.assertEqual(conflict_sheet["K9"].value, "Trường thử")
                self.assertEqual(conflict_sheet["P9"].value, "THCS")
                self.assertEqual(conflict_sheet["R9"].value, "2025")
                self.assertEqual(
                    tuple(
                        conflict_sheet[address].value
                        for address in (
                            "Q9",
                            "S9",
                            "T9",
                            "V9",
                            "W9",
                            "X9",
                            "Y9",
                            "Z9",
                            "AA9",
                            "AB9",
                        )
                    ),
                    (None,) * 10,
                )
            finally:
                missing.close()
                conflict.close()

    def test_approved_education_mapping_uses_only_first_member_row(self):
        education_values = (
            {
                0: {
                    "Q": "5",
                    "R": "5A1",
                    "T": "Trường giả lập Tiểu học",
                    "V": "MN",
                    "X": "2018-2019",
                },
                1: {
                    "Q": "3-4 tuổi",
                    "R": "M4**",
                    "T": "1982-1983",
                },
                2: {
                    "F": 1970,
                    "V": "THPT",
                    "X": "1998-1999",
                },
                4: {
                    "Q": "8",
                    "R": "7A1",
                    "T": "Trường giả lập THCS",
                },
                5: {
                    "Q": "6",
                    "R": "6A4",
                    "T": "Trường giả lập số 6",
                    "V": "TH",
                },
            },
            {
                6: {
                    "Q": "?",
                    "R": "CHUYỂN",
                    "T": "1982-1983",
                    "X": "2005-2006",
                }
            },
            {
                13: {
                    "Q": "12",
                    "R": "12A1",
                    "T": "Trường giả lập số 14",
                    "V": "THPT",
                    "X": "2020-2021",
                }
            },
        )
        counts = (6, 7, 14)
        households = tuple(
            make_household(
                f"SYNTHETIC-{count:02}",
                count,
                source_values_by_member=values,
            )
            for count, values in zip(counts, education_values, strict=True)
        )
        expected_label_values = {address: f"{2025 + index % 5}-{2026 + index % 5}" for index, address in enumerate(YEAR_LABEL_CELLS)}
        template_hash = hashlib.sha256(TEMPLATE.read_bytes()).digest()

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "education-mapping"
            report = export_households(
                households,
                TEMPLATE,
                output,
                MAPPING,
                settings_path=write_export_settings(
                    directory,
                    school_year="2026-2027",
                ),
            )
            self.assertEqual(report.success_count, 3)
            self.assertEqual(report.failure_count, 0)
            by_code = {item.household_code: item for item in report.successes}
            self.assertEqual(
                {code: item.member_count for code, item in by_code.items()},
                {
                    "SYNTHETIC-06": 6,
                    "SYNTHETIC-07": 7,
                    "SYNTHETIC-14": 14,
                },
            )

            warning_codes = {
                code: {warning.code for warning in item.warnings}
                for code, item in by_code.items()
            }
            self.assertIn(
                "B5_CLASS_VALUE_REQUIRES_REVIEW",
                warning_codes["SYNTHETIC-06"],
            )
            self.assertIn(
                "B5_GRADE_CLASS_CONFLICT",
                warning_codes["SYNTHETIC-06"],
            )
            self.assertIn(
                "B6_SCHOOL_NAME_LOOKS_LIKE_YEAR",
                warning_codes["SYNTHETIC-06"],
            )
            self.assertIn(
                "B7_GRADUATION_PAIR_INCOMPLETE",
                warning_codes["SYNTHETIC-06"],
            )
            self.assertIn(
                "B5_CLASS_VALUE_REQUIRES_REVIEW",
                warning_codes["SYNTHETIC-07"],
            )
            self.assertIn(
                "B5_GRADE_CLASS_NOT_COMPARABLE",
                warning_codes["SYNTHETIC-07"],
            )
            self.assertIn(
                "B6_SCHOOL_NAME_LOOKS_LIKE_YEAR",
                warning_codes["SYNTHETIC-07"],
            )
            self.assertIn(
                "B7_GRADUATION_PAIR_INCOMPLETE",
                warning_codes["SYNTHETIC-07"],
            )
            self.assertNotIn(
                "UNMAPPED_CLASS_NAME",
                warning_codes["SYNTHETIC-06"],
            )
            self.assertNotIn(
                "UNMAPPED_SCHOOL_NAME",
                warning_codes["SYNTHETIC-06"],
            )

            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.EnableEvents = False
            workbook = None
            try:
                checks = {
                    "SYNTHETIC-06": {
                        0: ("5A1", "Trường giả lập Tiểu học", "MN", "2018-2019"),
                        1: ("M4**", "1982-1983", None, None),
                        2: (None, None, "THPT", "1998-1999"),
                        3: (None, None, None, None),
                        4: ("7A1", "Trường giả lập THCS", None, None),
                        5: ("6A4", "Trường giả lập số 6", "TH", None),
                    },
                    "SYNTHETIC-07": {
                        6: ("CHUYỂN", "1982-1983", None, "2005-2006")
                    },
                    "SYNTHETIC-14": {
                        13: ("12A1", "Trường giả lập số 14", "THPT", "2020-2021")
                    },
                }
                for code, member_expectations in checks.items():
                    workbook = excel.Workbooks.Open(
                        str(by_code[code].path),
                        UpdateLinks=0,
                        ReadOnly=True,
                        AddToMru=False,
                        IgnoreReadOnlyRecommended=True,
                    )
                    sheet = workbook.Worksheets.Item("Phiếu điều tra 1 mặt")
                    for member_index, values in member_expectations.items():
                        page_index, slot_index = divmod(member_index, 7)
                        sheet = workbook.Worksheets.Item(page_index + 1)
                        start_row = (9, 14, 19, 24, 29, 34, 39)[slot_index]
                        for column, expected_value in zip(
                            ("H", "K", "P", "R"),
                            values,
                            strict=True,
                        ):
                            self.assertEqual(
                                sheet.Range(f"{column}{start_row}").Value2,
                                expected_value,
                                f"{code} member {member_index + 1} {column}",
                            )
                        self.assertEqual(
                            sheet.Range(f"B{start_row}").Value2,
                            f"Họ{member_index + 1} Tên{member_index + 1}",
                        )
                        for offset in range(1, 5):
                            for column in ("H", "K", "P", "R"):
                                self.assertIsNone(
                                    sheet.Range(
                                        f"{column}{start_row + offset}"
                                    ).Value2,
                                    f"Unexpected education value at "
                                    f"{column}{start_row + offset}",
                                )
                    self.assertEqual(
                        {
                            address: sheet.Range(address).Value2
                            for address in YEAR_LABEL_CELLS
                        },
                        expected_label_values,
                    )
                    self.assertEqual(
                        sheet.PageSetup.PrintArea,
                        "$A$1:$AG$87",
                    )
                    saved_workbook = load_workbook(
                        by_code[code].path,
                        read_only=False,
                    )
                    try:
                        self.assertEqual(
                            len(
                                saved_workbook[
                                    "Phiếu điều tra 1 mặt"
                                ].merged_cells.ranges
                            ),
                            442,
                        )
                    finally:
                        saved_workbook.close()
                    workbook.Close(SaveChanges=False)
                    workbook = None
            finally:
                if workbook is not None:
                    workbook.Close(SaveChanges=False)
                excel.Quit()
                pythoncom.CoUninitialize()

        self.assertEqual(
            hashlib.sha256(TEMPLATE.read_bytes()).digest(),
            template_hash,
        )

    def test_invalid_filename_characters_are_sanitized_and_existing_files_are_not_overwritten(self):
        first = make_household("A/B", 1)
        second = make_household("A:B", 1)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            target = output / "ho_A_B.xlsx"
            report = export_households(
                (first, second),
                TEMPLATE,
                output,
                MAPPING,
                settings_path=write_export_settings(directory),
            )

            self.assertEqual(report.success_count, 1)
            self.assertEqual(report.failure_count, 1)
            self.assertEqual(report.summary["success_count"], 1)
            self.assertEqual(report.summary["failure_count"], 1)
            self.assertTrue(target.is_file())
            original_result = target.read_bytes()
            self.assertIn("already exists", report.failures[0].error)

            repeated = export_households(
                (first,),
                TEMPLATE,
                output,
                MAPPING,
                settings_path=write_export_settings(directory),
            )

            self.assertEqual(repeated.success_count, 0)
            self.assertEqual(repeated.failure_count, 1)
            self.assertEqual(target.read_bytes(), original_result)

    def test_refuses_to_overwrite_a_formula_in_a_mapped_template_cell(self):
        import pythoncom
        import win32com.client

        household = make_household("FORMULA", 1)
        with tempfile.TemporaryDirectory() as directory:
            modified_template = Path(directory) / "formula-template.xls"
            shutil.copyfile(TEMPLATE, modified_template)
            pythoncom.CoInitialize()
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = None
            try:
                workbook = excel.Workbooks.Open(
                    str(modified_template), UpdateLinks=0, ReadOnly=False
                )
                cell = workbook.Worksheets.Item(1).Range("B9")
                cell.NumberFormat = "General"
                cell.Formula = "=1+1"
                workbook.Save()
                workbook.Close(SaveChanges=False)
                workbook = None
            finally:
                if workbook is not None:
                    workbook.Close(SaveChanges=False)
                excel.Quit()
                excel = None
                pythoncom.CoUninitialize()

            report = export_households(
                (household,),
                modified_template,
                Path(directory) / "output",
                MAPPING,
                settings_path=write_export_settings(directory),
            )

            self.assertEqual(report.success_count, 0)
            self.assertEqual(report.failure_count, 1)
            self.assertIn("formula", report.failures[0].error.lower())
            self.assertFalse((Path(directory) / "output" / "ho_FORMULA.xlsx").exists())

    def test_export_requires_a_configured_commune_instead_of_using_the_template_default(self):
        household = make_household("NO-COMMUNE", 1, commune_name="")

        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ExportConfigurationError):
                export_households(
                    (household,),
                    TEMPLATE,
                    Path(directory) / "output",
                    MAPPING,
                    settings_path=write_export_settings(directory),
                )
            self.assertFalse((Path(directory) / "output").exists())

    def test_export_rejects_missing_commune_before_creating_output(self):
        household = make_household("NO-CONFIG", 1)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "not-created"
            with self.assertRaisesRegex(
                ExportConfigurationError,
                "commune_name",
            ):
                export_households(
                    (household,),
                    TEMPLATE,
                    output,
                    MAPPING,
                )
            self.assertFalse(output.exists())

    def test_export_rejects_invalid_school_year_and_commune_mismatch(self):
        household = make_household("BAD-CONFIG", 1)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "not-created"
            with self.assertRaisesRegex(
                ExportConfigurationError,
                "school_year.*YYYY-YYYY",
            ):
                export_households(
                    (household,),
                    TEMPLATE,
                    output,
                    MAPPING,
                    settings_path=write_export_settings(
                        directory,
                        school_year="đồng ý",
                    ),
                )
            self.assertFalse(output.exists())

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "not-created"
            with self.assertRaisesRegex(
                ExportConfigurationError,
                "does not match configured commune_name",
            ):
                export_households(
                    (make_household("OTHER-COMMUNE", 1, commune_name="Xã Khác"),),
                    TEMPLATE,
                    output,
                    MAPPING,
                    settings_path=write_export_settings(directory),
                )
            self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
