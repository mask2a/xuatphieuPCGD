import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from analysis import analyze_sources
from reader import RawRecord, SourceLocation
from settings import Settings
from src.approved_mapping import address_text, resolve, values, FIELDS, DISABILITIES
from src.excel_exporter import export_households
from tests.test_excel_exporter import make_household, write_export_settings, ROOT, TEMPLATE, MAPPING


def full_household(code, count):
    household = make_household(code, count)
    headers = {f"C{i}": title for i, title in enumerate([v[0] for v in FIELDS.values()] + list(DISABILITIES))}
    raw_values = {column: ("x" if title in DISABILITIES else f"TEST {index}") for index, (column, title) in enumerate(headers.items())}
    members = []
    for index, member in enumerate(household.members):
        record = replace(member.source_record, headers=headers, values={**member.source_record.values, **raw_values})
        members.append(replace(member, source_record=record, fields={**member.fields, "female": "X" if index % 2 == 0 else "", "address": "Tổ kiểm thử"}))
    return replace(household, members=tuple(members))


class ApprovedMappingTests(unittest.TestCase):
    def test_address_handles_empty_values_and_whole_locality(self):
        for address, commune, expected in (
            (None, "Phường A", "Phường A"),
            ("12", "", "12"),
            ("12, phường a", "Phường A", "12, phường a"),
            ("12, Phường An", "Phường A", "12, Phường An, Phường A"),
            ("", "", ""),
        ):
            self.assertEqual(address_text(address, commune), expected)

    def test_headers_are_resolved_by_title_not_position_and_reject_duplicates(self):
        record = RawRecord({"ZZ": "nguyên văn"}, {}, SourceLocation("test", "test", 5), {"ZZ": "Ghi chú"})
        self.assertEqual(resolve(record, "Ghi chú"), ("nguyên văn", None))
        self.assertEqual(resolve(replace(record, headers={"ZZ": "Ghi chú", "A": "Ghi chú"}), "Ghi chú"), (None, "SOURCE_HEADER_UNRESOLVED"))
        self.assertEqual(resolve(replace(record, headers={}), "Ghi chú"), (None, "SOURCE_HEADER_UNRESOLVED"))

    def test_disability_combines_only_marked_titles_and_preserves_values(self):
        member = full_household("TEST", 1).members[0]
        result, warnings = values(member)
        self.assertFalse(warnings)
        self.assertEqual(result["AC"], "Vận động, Nghe, nói, Nhìn, Thần kinh, tâm thần, Trí tuệ, Học tập, Tự kỷ, Khác")
        self.assertEqual(result["S"], "TEST 0")
        self.assertEqual(result["AG"], "TEST 12")

    def test_merged_source_headers_and_missing_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "MauNhapLieu"
            for cell, value in {"B2": "Họ đệm", "C2": "Tên", "N2": "Số phiếu", "AB2": "Học xong", "AB3": "Lớp", "AC3": "Năm", "B5": "Test", "C5": "A", "N5": "1", "AC5": "2020-2021"}.items():
                sheet[cell] = value
            sheet.merge_cells("AB2:AC2")
            workbook.save(path)
            workbook.close()
            result = analyze_sources([path], Settings())
            self.assertEqual(resolve(result.read_result.records[0], "Học xong / Năm"), ("2020-2021", None))
            workbook = load_workbook(path)
            workbook.active["N2"] = "Sai tiêu đề"
            workbook.save(path)
            workbook.close()
            self.assertTrue(analyze_sources([path], Settings()).read_result.errors)

    def test_new_mappings_on_seven_and_eight_member_forms(self):
        with tempfile.TemporaryDirectory() as directory:
            households = [full_household(f"MAP-{count}", count) for count in (7, 8)]
            report = export_households(households, TEMPLATE, Path(directory) / "out", MAPPING, settings_path=write_export_settings(directory, school_year="2028-2029"))
            self.assertEqual(report.failure_count, 0, report.failures)
            self.assertEqual(report.success_count, 2)
            for item in report.successes:
                workbook = load_workbook(item.path)
                try:
                    sheet = workbook.worksheets[0]
                    self.assertEqual(sheet["A3"].value, "Xóm (2): Tổ kiểm thử, Xã Kiểm thử")
                    self.assertIn("☒ Thường trú", sheet["Z2"].value)
                    self.assertEqual(sheet["Z2"].number_format, "General")
                    self.assertIn("☐ Tạm trú", sheet["Z2"].value)
                    self.assertEqual(sheet["C12"].value, "☒")
                    self.assertEqual(sheet["C17"].value, "☐")
                    self.assertEqual(sheet["S9"].value, "TEST 0")
                    self.assertEqual(sheet["AG9"].value, "TEST 12")
                    for row in (9, 14, 19, 24, 29, 34, 39, 47):
                        self.assertEqual(sheet[f"F{row}"].value, "2025-2026")
                        self.assertIsNone(sheet[f"S{row + 1}"].value)
                    self.assertTrue(sheet.row_dimensions[47].hidden)
                    self.assertFalse(sheet.row_dimensions[39].hidden)
                    self.assertEqual(len(sheet.merged_cells.ranges), 442)
                finally:
                    workbook.close()
