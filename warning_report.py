"""Write the phase-two household and data-warning report to an XLSX workbook."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analysis import AnalysisResult


def write_warning_report(result: AnalysisResult, path: str | Path) -> Path:
    """Create a summary, household-count, and warning workbook."""
    report_path = Path(path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tong quan"
    household_sheet = workbook.create_sheet("So ho")
    warning_sheet = workbook.create_sheet("Canh bao")

    summary_rows = [
        ("Chỉ số", "Giá trị"),
        ("Tổng số tệp đọc", len(result.read_result.files_read)),
        ("Tổng số tệp lỗi", len(result.read_result.errors)),
        ("Tổng số tệp bỏ qua", len(result.read_result.skipped_files)),
        ("Tổng số bản ghi", len(result.members)),
        ("Tổng số hộ", len(result.grouping_result.households)),
        (
            "Số hộ có cảnh báo",
            sum(bool(household.warnings) for household in result.grouping_result.households),
        ),
        ("Xã cấu hình", result.settings.commune_name or ""),
        ("Năm học cấu hình", result.settings.school_year or ""),
    ]
    for row in summary_rows:
        summary.append(row)

    household_sheet.append(
        [
            "Mã hộ",
            "Xã",
            "Tệp nguồn",
            "Số thành viên",
            "Tên chủ hộ khai báo",
            "Tên chủ hộ đề xuất",
            "Điện thoại chủ hộ",
            "Diện cư trú",
            "Mã cảnh báo",
        ]
    )
    for household in result.grouping_result.households:
        household_sheet.append(
            [
                _display(household.household_code),
                household.commune_name,
                household.source_file,
                len(household.members),
                household.household_head_name,
                household.suggested_head_name,
                household.phone,
                household.residence,
                ", ".join(warning.code for warning in household.warnings),
            ]
        )

    warning_sheet.append(
        ["Mức độ", "Mã cảnh báo", "Mã hộ", "Tệp nguồn", "Dòng nguồn", "Nội dung"]
    )
    for household in result.grouping_result.households:
        for warning in household.warnings:
            warning_sheet.append(
                [
                    "Cảnh báo",
                    warning.code,
                    _display(household.household_code),
                    household.source_file,
                    _source_rows(household),
                    warning.message,
                ]
            )
    for error in result.read_result.errors:
        warning_sheet.append(
            ["Lỗi", "FILE_READ_ERROR", "", str(error.file), "", error.message]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_width = max(
                (
                    min(len(str(cell.value or "")), 60)
                    for cell in column_cells
                ),
                default=12,
            )
            sheet.column_dimensions[column_letter].width = max(12, max_width + 2)

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{report_path.stem}_", suffix=report_path.suffix, dir=report_path.parent
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, report_path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()
    return report_path


def _display(value: object) -> str:
    return "" if value is None else str(value)


def _source_rows(household) -> str:
    return ", ".join(
        str(member.source_record.source.row) for member in household.members
    )
