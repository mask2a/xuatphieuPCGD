"""Read-only, field-by-field audit of the six previously exported fixtures."""
import hashlib
import json
import shutil
import sys
import tempfile
from copy import copy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pythoncom
import win32com.client
from openpyxl import load_workbook

from src.excel_exporter import _verify_template_sheet, _template_signature
from src.template_mapper import load_template_mapping
from tests.test_approved_mapping import full_household

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "output/testing/approved_mapping_20261021_010939")
ROWS = (9, 14, 19, 24, 29, 34, 39, 47, 52, 57, 62, 67, 72, 77)
EXTRA_COLUMNS = ("S", "T", "V", "W", "X", "Y", "Z", "AA", "AB", "AD", "AE", "AF", "AG")
DISABILITY = "Vận động, Nghe, nói, Nhìn, Thần kinh, tâm thần, Trí tuệ, Học tập, Tự kỷ, Khác"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    template = ROOT / "maubaocao.xls"
    mapping = load_template_mapping(ROOT / "config/template_mapping.json")
    protected = [template, ROOT / "dulieumau.xls", ROOT / "config/approved_workbook_snapshot.json"]
    paths = [FOLDER / f"ho_SYNTHETIC-{count:02}.xlsx" for count in (1, 5, 7, 8, 14, 15)]
    before = {str(p.relative_to(ROOT)): digest(p) for p in protected + paths}
    report = {"mode": "read_only_existing_files", "hashes_before": before, "files": [], "failures": []}

    def check(condition, context):
        if not condition:
            report["failures"].append(context)

    pythoncom.CoInitialize()
    excel = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = False
        excel.AutomationSecurity = 3
        with tempfile.TemporaryDirectory() as directory:
            copied = Path(directory) / "template.xls"
            baseline_path = Path(directory) / "template.xlsx"
            shutil.copy2(template, copied)
            workbook = excel.Workbooks.Open(str(copied), UpdateLinks=0, ReadOnly=False, AddToMru=False)
            try:
                sheet = _verify_template_sheet(workbook, mapping)
                _template_signature(template, sheet, mapping)
                workbook.SaveAs(str(baseline_path), FileFormat=51, AddToMru=False)
            finally:
                workbook.Close(SaveChanges=False)
            baseline_book = load_workbook(baseline_path)
            baseline = baseline_book.worksheets[0]
            try:
                for count, path in zip((1, 5, 7, 8, 14, 15), paths):
                    household = full_household(f"SYNTHETIC-{count:02}", count)
                    output = load_workbook(path)
                    checks = 0
                    try:
                        check(len(output.worksheets) == (2 if count > 14 else 1), f"{path.name}: sheet count")
                        for page_index, sheet in enumerate(output.worksheets):
                            label = f"{path.name}/{sheet.title}"
                            members = household.members[page_index * 14:(page_index + 1) * 14]
                            expected = {
                                "A1": "Xã (1): Xã Kiểm thử", "A3": "Xóm (2): Tổ kiểm thử, Xã Kiểm thử",
                                "H3": "Họ và tên chủ hộ (3): Chủ Hộ", "Z1": f"Số phiếu (4): SYNTHETIC-{count:02}",
                                "Z3": "Điện thoại (6): 0912345678", "Z2": "Diện cư trú (5): X Thường trú     Tạm trú",
                            }
                            commune = str(baseline["A1"].value).split(":", 1)[1].strip()
                            expected["AA82"] = baseline["AA82"].value.replace(commune, household.commune_name)
                            for slot, row in enumerate(ROWS):
                                for offset in range(5):
                                    expected[f"F{row + offset}"] = f"{2025 + offset}-{2026 + offset}"
                                if slot < len(members):
                                    index = page_index * 14 + slot
                                    expected.update({
                                        f"A{row}": index + 1, f"B{row}": f"Họ{index + 1} Tên{index + 1}",
                                        f"B{row + 1}": "QH với chủ hộ: " + ("Chủ hộ" if index == 0 else "Thành viên"),
                                        f"D{row + 2}": f"{index + 1:02}/{index + 2:02}/{2010 + index}",
                                        f"C{row + 3}": "X" if index % 2 == 0 else "",
                                        f"D{row + 3}": "DT: Kinh", f"B{row + 4}": "Cha, mẹ, Ng đỡ đầu: Cha mẹ",
                                        f"AC{row}": DISABILITY,
                                    })
                                    for number, column in enumerate(EXTRA_COLUMNS):
                                        expected[f"{column}{row}"] = f"TEST {number}"
                            for row in range(1, 88):
                                check(sheet.row_dimensions[row].height == baseline.row_dimensions[row].height, f"{label}: row height {row}")
                                for column in range(1, 34):
                                    actual = sheet.cell(row, column)
                                    original = baseline.cell(row, column)
                                    address = actual.coordinate
                                    value = expected.get(address, original.value)
                                    check((actual.value or "") == (value or ""), f"{label}: value {address}")
                                    font = copy(original.font)
                                    if address.startswith("C") and any(address == f"C{r + 3}" and expected.get(address) == "X" for r in ROWS):
                                        font.name = "Arial"
                                        font.family = 2
                                        font.charset = None
                                    check(actual.font == font, f"{label}: font {address}")
                                    for attribute in ("border", "fill", "alignment", "number_format", "protection"):
                                        check(copy(getattr(actual, attribute)) == copy(getattr(original, attribute)), f"{label}: {attribute} {address}")
                                    checks += 1
                            check({str(area) for area in sheet.merged_cells.ranges} == {str(area) for area in baseline.merged_cells.ranges}, f"{label}: merged ranges")
                            for column, dimension in baseline.column_dimensions.items():
                                check(sheet.column_dimensions[column].width == dimension.width, f"{label}: column width {column}")
                            for slot, row in enumerate(ROWS):
                                hidden = len(members) <= 7 and slot >= 7
                                check(bool(sheet.row_dimensions[row].hidden) == hidden, f"{label}: visible block {slot + 1}")
                            check(sheet.page_margins == baseline.page_margins, f"{label}: margins")
                            for column in ("F", "J", "N", "R", "W"):
                                check(sheet[f"{column}82"].value == baseline[f"{column}82"].value, f"{label}: signature year {column}")
                        check(not output._external_links, f"{path.name}: external links")
                    finally:
                        output.close()
                    workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True, AddToMru=False)
                    try:
                        pages = []
                        for sheet in workbook.Worksheets:
                            sheet.Activate()
                            pages.append(int(excel.ExecuteExcel4Macro("GET.DOCUMENT(50)")))
                            check(sheet.PageSetup.PrintArea == "$A$1:$AG$87", f"{path.name}: print area")
                            check(sheet.PageSetup.PaperSize == 9 and sheet.PageSetup.Orientation == 2, f"{path.name}: paper size/orientation")
                        check(pages == ([1, 1] if count > 14 else [1]), f"{path.name}: printed page count")
                    finally:
                        workbook.Close(SaveChanges=False)
                    report["files"].append({"file": path.name, "members": count, "cells_compared": checks, "pages": pages})
            finally:
                baseline_book.close()
    finally:
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()
    after = {str(p.relative_to(ROOT)): digest(p) for p in protected + paths}
    check(before == after, "Protected originals or existing sample files changed")
    report["files_unchanged"] = before == after
    report_path = ROOT / "docs" / ("EXISTING_SAMPLES_AUDIT.json" if len(sys.argv) == 1 else "REVIEWED_SAMPLES_AUDIT.json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"files": report["files"], "failure_count": len(report["failures"]), "failures": report["failures"][:40], "unchanged": report["files_unchanged"]}, ensure_ascii=True))
    return bool(report["failures"])


if __name__ == "__main__":
    raise SystemExit(main())
